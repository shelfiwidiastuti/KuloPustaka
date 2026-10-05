from flask import Flask, render_template, request, redirect, url_for, session, flash, jsonify
import mysql.connector
from mysql.connector import Error
from functools import wraps
from datetime import date, datetime
import os
from openpyxl import load_workbook
from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    session,
    flash,
    jsonify
)

import mysql.connector
from mysql.connector import Error

from functools import wraps
from datetime import date, datetime
from werkzeug.security import generate_password_hash, check_password_hash

import os
app = Flask(__name__)

app.secret_key = os.environ.get(
    "SECRET_KEY",
    "kulopustaka-secret-key-change-me"
)


# =========================================================
# DATABASE
# =========================================================

DB_CONFIG = {
    "host": os.environ.get("DB_HOST", "localhost"),
    "user": os.environ.get("DB_USER", "root"),
    "password": os.environ.get("DB_PASSWORD", ""),
    "database": os.environ.get("DB_NAME", "kulopustaka")
}


def get_db():
    return mysql.connector.connect(**DB_CONFIG)

def ensure_admin_account():
    """
    Membuat tabel admin_users jika belum ada
    dan membuat akun admin pertama kali.
    """

    db = get_db()
    cur = db.cursor()

    cur.execute("""
        CREATE TABLE IF NOT EXISTS admin_users (
            id INT AUTO_INCREMENT PRIMARY KEY,
            username VARCHAR(50) NOT NULL UNIQUE,
            password_hash VARCHAR(255) NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                ON UPDATE CURRENT_TIMESTAMP
        )
    """)

    # Cek apakah akun admin sudah ada
    cur.execute("""
        SELECT id
        FROM admin_users
        WHERE username = %s
        LIMIT 1
    """, ("admin",))

    admin = cur.fetchone()

    # Jika belum ada, buat akun admin pertama
    if not admin:
        password_hash = generate_password_hash("KuloPustaka123")

        cur.execute("""
            INSERT INTO admin_users
            (
                username,
                password_hash
            )
            VALUES (%s, %s)
        """, (
            "admin",
            password_hash
        ))

    db.commit()

    cur.close()
    db.close()
# =========================================================
# LOGIN REQUIRED
# =========================================================

def login_required(view):

    @wraps(view)
    def wrapped(*args, **kwargs):

        if not session.get("logged_in"):
            return redirect(url_for("login"))

        return view(*args, **kwargs)

    return wrapped


# =========================================================
# GLOBAL VARIABLES UNTUK TEMPLATE
# =========================================================

@app.context_processor
def inject_globals():

    return {
        "current_year": datetime.now().year,
        "today": date.today(),
        "current_date": date.today()
    }


# =========================================================
# INDEX
# =========================================================

@app.route("/")
def index():

    if session.get("logged_in"):
        return redirect(url_for("dashboard"))

    return redirect(url_for("login"))


# =========================================================
# LOGIN
# =========================================================
@app.route("/login", methods=["GET", "POST"])
def login():

    if request.method == "POST":

        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")

        conn = get_db()
        cursor = conn.cursor(dictionary=True)

        cursor.execute(
            """
            SELECT id, username, password_hash
            FROM admin_users
            WHERE username = %s
            """,
            (username,)
        )

        user = cursor.fetchone()

        cursor.close()
        conn.close()

        if user and check_password_hash(
            user["password_hash"],
            password
        ):

            # SIMPAN DATA ADMIN KE SESSION
            session["logged_in"] = True
            session["admin_id"] = user["id"]
            session["username"] = user["username"]

            return redirect(url_for("dashboard"))

        flash("Username atau kata sandi salah.", "error")

    return render_template("login.html")

# =========================================================
# LOGOUT
# =========================================================

@app.route("/logout")
def logout():

    session.clear()

    return redirect(url_for("login"))


# =========================================================
# DASHBOARD
# =========================================================
@app.route("/ubah-kata-sandi", methods=["GET", "POST"])
@login_required
def change_password():

    if request.method == "POST":

        old_password = request.form.get("old_password", "")
        new_password = request.form.get("new_password", "")
        confirm_password = request.form.get("confirm_password", "")

        # Validasi
        if not old_password or not new_password or not confirm_password:
            flash("Semua kolom kata sandi wajib diisi.", "error")
            return redirect(url_for("change_password"))

        if len(new_password) < 8:
            flash(
                "Kata sandi baru minimal 8 karakter.",
                "error"
            )
            return redirect(url_for("change_password"))

        if new_password != confirm_password:
            flash(
                "Konfirmasi kata sandi baru tidak cocok.",
                "error"
            )
            return redirect(url_for("change_password"))

        # Ambil data admin yang sedang login
        db = get_db()
        cur = db.cursor(dictionary=True)

        cur.execute("""
            SELECT id, password_hash
            FROM admin_users
            WHERE id = %s
            LIMIT 1
        """, (
            session["admin_id"],
        ))

        admin = cur.fetchone()

        if not admin:
            cur.close()
            db.close()

            session.clear()

            flash(
                "Akun admin tidak ditemukan. Silakan login kembali.",
                "error"
            )

            return redirect(url_for("login"))

        # Cek password lama
        if not check_password_hash(
            admin["password_hash"],
            old_password
        ):
            cur.close()
            db.close()

            flash(
                "Kata sandi lama salah.",
                "error"
            )

            return redirect(url_for("change_password"))

        # Password baru tidak boleh sama
        if check_password_hash(
            admin["password_hash"],
            new_password
        ):
            cur.close()
            db.close()

            flash(
                "Kata sandi baru harus berbeda dari kata sandi lama.",
                "error"
            )

            return redirect(url_for("change_password"))

        # Hash password baru
        new_password_hash = generate_password_hash(
            new_password
        )

        # Simpan password baru
        cur.execute("""
            UPDATE admin_users
            SET password_hash = %s
            WHERE id = %s
        """, (
            new_password_hash,
            session["admin_id"]
        ))

        db.commit()

        cur.close()
        db.close()

        flash(
            "Kata sandi berhasil diubah.",
            "success"
        )

        return redirect(url_for("dashboard"))

    return render_template("change_password.html")

@app.route("/dashboard")
@login_required
def dashboard():

    db = get_db()
    cur = db.cursor(dictionary=True)

    # -----------------------------------------------------
    # STATISTIK BUKU
    # -----------------------------------------------------

    cur.execute("""
        SELECT COUNT(*) AS total
        FROM books
    """)

    total_books = cur.fetchone()["total"]


    cur.execute("""
        SELECT COALESCE(SUM(stock), 0) AS available
        FROM books
    """)

    available_books = cur.fetchone()["available"]


    # -----------------------------------------------------
    # STATISTIK PEMINJAMAN
    # -----------------------------------------------------

    cur.execute("""
        SELECT COUNT(*) AS total
        FROM loans
        WHERE status = 'Dipinjam'
    """)

    borrowed_books = cur.fetchone()["total"]


    cur.execute("""
        SELECT COUNT(*) AS total
        FROM loans
        WHERE status = 'Dipinjam'
        AND due_date < CURDATE()
    """)

    overdue_books = cur.fetchone()["total"]


    # -----------------------------------------------------
    # STATISTIK PEMINJAM
    # -----------------------------------------------------

    cur.execute("""
        SELECT COUNT(*) AS total
        FROM borrowers
    """)

    total_borrowers = cur.fetchone()["total"]


    # -----------------------------------------------------
    # BUKU TAMU - JUMLAH HARI INI
    # -----------------------------------------------------

    cur.execute("""
        SELECT COUNT(*) AS total
        FROM guestbook
        WHERE visit_date = CURDATE()
    """)

    today_visitors = cur.fetchone()["total"]


    # -----------------------------------------------------
    # PEMINJAMAN TERBARU
    # -----------------------------------------------------

    cur.execute("""
        SELECT
            l.id,
            b.title,
            br.name,
            br.member_code,
            l.loan_date,
            l.due_date,
            l.return_date,
            l.status
        FROM loans l
        JOIN books b
            ON b.id = l.book_id
        JOIN borrowers br
            ON br.id = l.borrower_id
        ORDER BY l.id DESC
        LIMIT 8
    """)

    recent = cur.fetchall()


    # -----------------------------------------------------
    # BUKU TAMU TERBARU
    # -----------------------------------------------------

    cur.execute("""
        SELECT
            id,
            name,
            origin,
            purpose,
            visit_date,
            entry_time,
            exit_time,
            note
        FROM guestbook
        ORDER BY visit_date DESC, entry_time DESC, id DESC
        LIMIT 5
    """)

    recent_guests = cur.fetchall()


    cur.close()
    db.close()


    # -----------------------------------------------------
    # KIRIM DATA KE DASHBOARD
    # -----------------------------------------------------

    stats = {
        "total_books": total_books,
        "available_books": available_books,
        "borrowed_books": borrowed_books,
        "overdue_books": overdue_books,
        "total_borrowers": total_borrowers,
        "today_visitors": today_visitors
    }


    return render_template(
        "dashboard.html",
        stats=stats,
        recent=recent,
        recent_guests=recent_guests
    )


# =========================================================
# DATA BUKU
# =========================================================

@app.route("/buku")
@login_required
def books():

    q = request.args.get("q", "").strip()

    db = get_db()
    cur = db.cursor(dictionary=True)

    if q:

        cur.execute("""
            SELECT *
            FROM books
            WHERE title LIKE %s
            OR author LIKE %s
            OR category LIKE %s
            ORDER BY id DESC
        """, (
            f"%{q}%",
            f"%{q}%",
            f"%{q}%"
        ))

    else:

        cur.execute("""
            SELECT *
            FROM books
            ORDER BY id DESC
        """)

    data = cur.fetchall()

    cur.close()
    db.close()

    return render_template(
        "books.html",
        books=data,
        q=q
    )


# =========================================================
# TAMBAH BUKU
# =========================================================
@app.route("/buku/tambah", methods=["GET", "POST"])
@login_required
def add_book():

    if request.method == "POST":

        code = request.form.get("code", "").strip() or None
        title = request.form.get("title", "").strip()
        author = request.form.get("author", "").strip() or None
        category = request.form.get("category", "").strip() or None
        publisher = request.form.get("publisher", "").strip() or None
        year = request.form.get("year") or None
        stock = request.form.get("stock") or 0
        isbn = request.form.get("isbn", "").strip() or None

        # Judul wajib
        if not title:
            flash("Judul buku wajib diisi.", "error")
            return redirect(url_for("books"))

        conn = None
        cur = None

        try:
            conn = get_db()
            cur = conn.cursor(dictionary=True)

            # Cek kode hanya jika kode diisi
            if code:
                cur.execute("""
                    SELECT id
                    FROM books
                    WHERE code = %s
                    LIMIT 1
                """, (code,))

                existing_book = cur.fetchone()

                if existing_book:
                    flash(
                        f"Kode buku '{code}' sudah terdaftar.",
                        "error"
                    )
                    return redirect(url_for("books"))

            # Tahun
            if year:
                try:
                    year = int(year)
                except (ValueError, TypeError):
                    year = None

            # Stok
            try:
                stock = int(stock)
                if stock < 0:
                    stock = 0
            except (ValueError, TypeError):
                stock = 0

            cur.execute("""
                INSERT INTO books
                (
                    code,
                    title,
                    author,
                    category,
                    publisher,
                    year,
                    stock,
                    isbn
                )
                VALUES
                (%s, %s, %s, %s, %s, %s, %s, %s)
            """, (
                code,
                title,
                author,
                category,
                publisher,
                year,
                stock,
                isbn
            ))

            conn.commit()

            flash("Buku berhasil ditambahkan.", "success")
            return redirect(url_for("books"))

        except Exception as e:

            if conn:
                conn.rollback()

            flash(f"Gagal menambahkan buku: {str(e)}", "error")
            return redirect(url_for("books"))

        finally:

            if cur:
                cur.close()

            if conn:
                conn.close()

    return render_template("books.html")
# =========================================================
# IMPORT DATA BUKU DARI EXCEL
# =========================================================

@app.route("/buku/import-excel", methods=["POST"])
@login_required
def import_books_excel():

    # -----------------------------------------------------
    # CEK FILE
    # -----------------------------------------------------

    if "file" not in request.files:

        flash(
            "File Excel belum dipilih.",
            "error"
        )

        return redirect(url_for("books"))


    file = request.files["file"]


    if file.filename == "":

        flash(
            "File Excel belum dipilih.",
            "error"
        )

        return redirect(url_for("books"))


    if not file.filename.lower().endswith(".xlsx"):

        flash(
            "File harus berformat .xlsx.",
            "error"
        )

        return redirect(url_for("books"))


    conn = None
    cur = None


    try:

        # -------------------------------------------------
        # BUKA FILE EXCEL
        # -------------------------------------------------

        workbook = load_workbook(
            file,
            data_only=True
        )


        # -------------------------------------------------
        # PILIH SHEET DATA BUKU
        # -------------------------------------------------

        if "Data Buku" not in workbook.sheetnames:

            flash(
                "Sheet 'Data Buku' tidak ditemukan di file Excel.",
                "error"
            )

            return redirect(url_for("books"))


        sheet = workbook["Data Buku"]


        # -------------------------------------------------
        # KONEKSI DATABASE
        # -------------------------------------------------

        conn = get_db()

        cur = conn.cursor(
            dictionary=True
        )


        berhasil = 0
        dilewati = 0


        # -------------------------------------------------
        # BACA DATA MULAI BARIS 2
        # -------------------------------------------------

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True
        ):

            # ---------------------------------------------
            # LEWATI BARIS BENAR-BENAR KOSONG
            # ---------------------------------------------

            if not row or all(
                value is None
                for value in row
            ):

                continue


            # ---------------------------------------------
            # SESUAI URUTAN KOLOM EXCEL
            #
            # A = No
            # B = Kode Buku
            # C = Judul Buku
            # D = Kategori
            # E = Penulis
            # F = Penerbit
            # G = Tahun
            # H = Stok
            # I = ISBN
            # ---------------------------------------------

            code = (
                row[1]
                if len(row) > 1
                else None
            )

            title = (
                row[2]
                if len(row) > 2
                else None
            )

            category = (
                row[3]
                if len(row) > 3
                else None
            )

            author = (
                row[4]
                if len(row) > 4
                else None
            )

            publisher = (
                row[5]
                if len(row) > 5
                else None
            )

            year = (
                row[6]
                if len(row) > 6
                else None
            )

            stock = (
                row[7]
                if len(row) > 7
                else None
            )

            isbn = (
                row[8]
                if len(row) > 8
                else None
            )


            # ---------------------------------------------
            # BERSIHKAN DATA
            # ---------------------------------------------

            if code is not None:

                code = str(code).strip()

                if code == "":
                    code = None


            if title is not None:

                title = str(title).strip()

            else:

                title = ""


            if category is not None:

                category = str(category).strip()

                if category == "":
                    category = None


            if author is not None:

                author = str(author).strip()

                if author == "":
                    author = None


            if publisher is not None:

                publisher = str(publisher).strip()

                if publisher == "":
                    publisher = None


            if isbn is not None:

                isbn = str(isbn).strip()

                if isbn == "":
                    isbn = None


            # ---------------------------------------------
            # JUDUL WAJIB
            # ---------------------------------------------

            if not title:

                dilewati += 1

                continue


            # ---------------------------------------------
            # TAHUN BOLEH KOSONG
            # ---------------------------------------------

            if year is not None:

                try:

                    if str(year).strip() != "":

                        year = int(
                            float(year)
                        )

                    else:

                        year = None

                except (
                    ValueError,
                    TypeError
                ):

                    year = None

            else:

                year = None


            # ---------------------------------------------
            # STOK BOLEH KOSONG
            # KOSONG = 0
            # ---------------------------------------------

            if stock is not None:

                try:

                    if str(stock).strip() != "":

                        stock = int(
                            float(stock)
                        )

                        if stock < 0:
                            stock = 0

                    else:

                        stock = 0

                except (
                    ValueError,
                    TypeError
                ):

                    stock = 0

            else:

                stock = 0


            # ---------------------------------------------
            # CEK KODE BUKU
            # HANYA JIKA DIISI
            # ---------------------------------------------

            if code:

                cur.execute("""
                    SELECT id
                    FROM books
                    WHERE code = %s
                    LIMIT 1
                """, (
                    code,
                ))


                existing_book = cur.fetchone()


                if existing_book:

                    dilewati += 1

                    continue


            # ---------------------------------------------
            # SIMPAN KE DATABASE
            # ---------------------------------------------

            cur.execute("""
                INSERT INTO books
                (
                    code,
                    title,
                    author,
                    category,
                    publisher,
                    year,
                    stock,
                    isbn
                )
                VALUES
                (
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s,
                    %s
                )
            """, (
                code,
                title,
                author,
                category,
                publisher,
                year,
                stock,
                isbn
            ))


            berhasil += 1


        # -------------------------------------------------
        # SIMPAN PERUBAHAN
        # -------------------------------------------------

        conn.commit()


        # -------------------------------------------------
        # PESAN HASIL
        # -------------------------------------------------

        flash(
            f"Import selesai. "
            f"Berhasil: {berhasil} buku. "
            f"Dilewati: {dilewati} baris.",
            "success"
        )


        return redirect(
            url_for("books")
        )


    except Exception as e:

        # -------------------------------------------------
        # JIKA ERROR
        # -------------------------------------------------

        if conn:

            conn.rollback()


        flash(
            f"Gagal mengimport Excel: {str(e)}",
            "error"
        )


        return redirect(
            url_for("books")
        )


    finally:

        if cur:
            cur.close()


        if conn:
            conn.close()
# =========================================================
# TEST BACA EXCEL
# =========================================================

@app.route("/test-excel", methods=["POST"])
@login_required
def test_excel():

    if "file" not in request.files:
        return "File tidak ditemukan"

    file = request.files["file"]

    if file.filename == "":
        return "File belum dipilih"

    workbook = load_workbook(file, data_only=False)

    print("\n========================================")
    print("TEST MEMBACA EXCEL")
    print("========================================")

    print("SHEET:")
    for sheet_name in workbook.sheetnames:
        print("-", sheet_name)

    sheet = workbook["Data Buku"]

    print("\nNILAI BARIS 2:")
    print(list(sheet.iter_rows(
        min_row=2,
        max_row=2,
        values_only=True
    )))

    print("\nCELL SATU-SATU:")

    for cell in ["A2", "B2", "C2", "D2", "E2", "F2", "G2", "H2", "I2"]:
        print(cell, "=", sheet[cell].value)

    print("========================================\n")

    return "Silakan lihat hasilnya di terminal Flask."
# =========================================================
# EDIT BUKU
# =========================================================

@app.route("/buku/edit/<int:id>", methods=["POST"])
@login_required
def edit_book(id):

    data = request.form

    db = get_db()
    cur = db.cursor()

    cur.execute("""
        UPDATE books
        SET
            code=%s,
            title=%s,
            author=%s,
            publisher=%s,
            year=%s,
            category=%s,
            stock=%s
        WHERE id=%s
    """, (
        data["code"],
        data["title"],
        data["author"],
        data["publisher"],
        data["year"] or None,
        data["category"],
        data["stock"],
        id
    ))

    db.commit()

    cur.close()
    db.close()

    flash("Data buku berhasil diperbarui.", "success")

    return redirect(url_for("books"))


# =========================================================
# HAPUS BUKU
# =========================================================

@app.route("/buku/hapus/<int:id>", methods=["POST"])
@login_required
def delete_book(id):

    db = get_db()
    cur = db.cursor()

    cur.execute(
        "DELETE FROM books WHERE id=%s",
        (id,)
    )

    db.commit()

    cur.close()
    db.close()

    flash("Buku berhasil dihapus.", "success")

    return redirect(url_for("books"))


# =========================================================
# DATA PEMINJAM
# =========================================================

@app.route("/peminjam")
@login_required
def borrowers():

    db = get_db()
    cur = db.cursor(dictionary=True)

    cur.execute("""
        SELECT
            br.*,
            COUNT(l.id) AS active_loans
        FROM borrowers br
        LEFT JOIN loans l
            ON br.id = l.borrower_id
            AND l.status = 'Dipinjam'
        GROUP BY br.id
        ORDER BY br.id DESC
    """)

    data = cur.fetchall()

    cur.close()
    db.close()

    return render_template(
        "borrowers.html",
        borrowers=data
    )


# =========================================================
# TAMBAH PEMINJAM
# =========================================================
@app.route("/peminjam/tambah", methods=["POST"])
@login_required
def add_borrower():

    data = request.form

    db = get_db()
    cur = db.cursor()

    # =========================================================
    # CEK KODE PEMINJAM
    # =========================================================

    cur.execute("""
        SELECT id
        FROM borrowers
        WHERE member_code = %s
    """, (
        data["member_code"],
    ))

    existing_borrower = cur.fetchone()

    if existing_borrower:

        cur.close()
        db.close()

        flash(
            f"Kode peminjam '{data['member_code']}' sudah terdaftar.",
            "error"
        )

        return redirect(url_for("borrowers"))

    # =========================================================
    # TAMBAH PEMINJAM
    # =========================================================

    cur.execute("""
        INSERT INTO borrowers
        (
            member_code,
            name,
            phone,
            address
        )
        VALUES (%s,%s,%s,%s)
    """, (
        data["member_code"],
        data["name"],
        data["phone"],
        data["address"]
    ))

    db.commit()

    cur.close()
    db.close()

    flash(
        "Peminjam berhasil ditambahkan.",
        "success"
    )

    return redirect(url_for("borrowers"))


# =========================================================
# HAPUS PEMINJAM
# =========================================================

@app.route("/peminjam/hapus/<int:id>", methods=["POST"])
@login_required
def delete_borrower(id):

    db = get_db()
    cur = db.cursor()

    cur.execute(
        "SELECT COUNT(*) FROM loans WHERE borrower_id=%s",
        (id,)
    )

    if cur.fetchone()[0]:

        flash(
            "Peminjam tidak dapat dihapus karena memiliki riwayat transaksi.",
            "danger"
        )

    else:

        cur.execute(
            "DELETE FROM borrowers WHERE id=%s",
            (id,)
        )

        db.commit()

        flash(
            "Peminjam berhasil dihapus.",
            "success"
        )

    cur.close()
    db.close()

    return redirect(url_for("borrowers"))


# =========================================================
# PEMINJAMAN
# =========================================================

@app.route("/peminjaman")
@login_required
def loans():

    db = get_db()
    cur = db.cursor(dictionary=True)

    cur.execute("""
        SELECT
            l.*,
            b.title,
            br.name,
            br.member_code
        FROM loans l
        JOIN books b
            ON b.id = l.book_id
        JOIN borrowers br
            ON br.id = l.borrower_id
        ORDER BY l.id DESC
    """)

    data = cur.fetchall()


    cur.execute("""
        SELECT
            id,
            title,
            stock
        FROM books
        WHERE stock > 0
        ORDER BY title
    """)

    books_data = cur.fetchall()


    cur.execute("""
        SELECT
            id,
            name,
            member_code
        FROM borrowers
        ORDER BY name
    """)

    borrowers_data = cur.fetchall()


    cur.close()
    db.close()

    return render_template(
        "loans.html",
        loans=data,
        books=books_data,
        borrowers=borrowers_data
    )


# =========================================================
# TAMBAH PEMINJAMAN
# =========================================================

@app.route("/peminjaman/tambah", methods=["POST"])
@login_required
def add_loan():

    data = request.form

    db = get_db()
    cur = db.cursor(dictionary=True)

    cur.execute(
        "SELECT stock FROM books WHERE id=%s FOR UPDATE",
        (data["book_id"],)
    )

    book = cur.fetchone()


    if not book or book["stock"] <= 0:

        db.rollback()

        cur.close()
        db.close()

        flash(
            "Stok buku tidak tersedia.",
            "danger"
        )

        return redirect(url_for("loans"))


    cur.execute("""
        INSERT INTO loans
        (
            book_id,
            borrower_id,
            loan_date,
            due_date,
            status
        )
        VALUES (%s,%s,%s,%s,'Dipinjam')
    """, (
        data["book_id"],
        data["borrower_id"],
        data["loan_date"],
        data["due_date"]
    ))


    cur.execute(
        "UPDATE books SET stock=stock-1 WHERE id=%s",
        (data["book_id"],)
    )


    db.commit()

    cur.close()
    db.close()

    flash(
        "Peminjaman berhasil dicatat.",
        "success"
    )

    return redirect(url_for("loans"))


# =========================================================
# PENGEMBALIAN
# =========================================================

@app.route("/pengembalian/<int:id>", methods=["POST"])
@login_required
def return_loan(id):

    db = get_db()
    cur = db.cursor(dictionary=True)

    cur.execute("""
        SELECT *
        FROM loans
        WHERE id=%s
        AND status='Dipinjam'
    """, (id,))

    loan = cur.fetchone()


    if not loan:

        flash(
            "Transaksi tidak ditemukan atau sudah dikembalikan.",
            "danger"
        )

    else:

        cur.execute("""
            UPDATE loans
            SET
                return_date=%s,
                status='Dikembalikan'
            WHERE id=%s
        """, (
            date.today(),
            id
        ))


        cur.execute(
            "UPDATE books SET stock=stock+1 WHERE id=%s",
            (loan["book_id"],)
        )


        db.commit()

        flash(
            "Buku berhasil dikembalikan.",
            "success"
        )


    cur.close()
    db.close()

    return redirect(url_for("loans"))


# =========================================================
# RIWAYAT
# =========================================================

@app.route("/riwayat")
@login_required
def history():

    db = get_db()
    cur = db.cursor(dictionary=True)

    cur.execute("""
        SELECT
            l.*,
            b.title,
            br.name,
            br.member_code
        FROM loans l
        JOIN books b
            ON b.id = l.book_id
        JOIN borrowers br
            ON br.id = l.borrower_id
        ORDER BY l.id DESC
    """)

    data = cur.fetchall()

    cur.close()
    db.close()

    return render_template(
        "history.html",
        loans=data
    )


# =========================================================
# LAPORAN
# =========================================================

@app.route("/laporan")
@login_required
def reports():

    db = get_db()
    cur = db.cursor(dictionary=True)

    cur.execute("""
        SELECT
            DATE_FORMAT(loan_date,'%Y-%m') AS month,
            COUNT(*) AS total
        FROM loans
        GROUP BY month
        ORDER BY month DESC
        LIMIT 12
    """)

    monthly = cur.fetchall()


    cur.execute("""
        SELECT
            b.title,
            COUNT(*) AS total
        FROM loans l
        JOIN books b
            ON b.id = l.book_id
        GROUP BY b.id
        ORDER BY total DESC
        LIMIT 10
    """)

    popular = cur.fetchall()


    cur.close()
    db.close()

    return render_template(
        "reports.html",
        monthly=monthly,
        popular=popular
    )


# =========================================================
# BUKU TAMU
# =========================================================

@app.route("/buku-tamu")
@login_required
def guestbook():

    q = request.args.get("q", "").strip()
    visit_date = request.args.get("visit_date", "").strip()

    db = get_db()
    cur = db.cursor(dictionary=True)


    query = """
        SELECT *
        FROM guestbook
        WHERE 1=1
    """

    params = []


    if q:

        query += """
            AND (
                name LIKE %s
                OR origin LIKE %s
                OR purpose LIKE %s
            )
        """

        keyword = f"%{q}%"

        params.extend([
            keyword,
            keyword,
            keyword
        ])


    if visit_date:

        query += """
            AND visit_date = %s
        """

        params.append(visit_date)


    query += """
        ORDER BY
            visit_date DESC,
            entry_time DESC,
            id DESC
    """


    cur.execute(
        query,
        tuple(params)
    )

    data = cur.fetchall()


    # -----------------------------------------------------
    # TAMU HARI INI
    # -----------------------------------------------------

    cur.execute("""
        SELECT COUNT(*) AS total
        FROM guestbook
        WHERE visit_date = CURDATE()
    """)

    today_visitors = cur.fetchone()["total"]


    # -----------------------------------------------------
    # TOTAL TAMU
    # -----------------------------------------------------

    cur.execute("""
        SELECT COUNT(*) AS total
        FROM guestbook
    """)

    total_visitors = cur.fetchone()["total"]


    cur.close()
    db.close()


    return render_template(
        "guestbook.html",
        guests=data,
        q=q,
        visit_date=visit_date,
        today_visitors=today_visitors,
        total_visitors=total_visitors
    )


# =========================================================
# TAMBAH BUKU TAMU
# =========================================================
@app.route("/buku-tamu/tambah", methods=["POST"])
@login_required
def add_guest():

    data = request.form

    db = get_db()
    cur = db.cursor()

    cur.execute("""
        INSERT INTO guestbook
        (
            name,
            origin,
            purpose,
            visit_date,
            entry_time,
            exit_time,
            note,
            paraf
        )
        VALUES (%s,%s,%s,%s,%s,%s,%s,%s)
    """, (
        data["name"],
        data["origin"],
        data["purpose"],
        data["visit_date"],
        data["entry_time"],
        data.get("exit_time") or None,
        data.get("note", ""),
        data.get("paraf", "")
    ))

    db.commit()

    cur.close()
    db.close()

    flash(
        "Data tamu berhasil ditambahkan.",
        "success"
    )

    return redirect(url_for("guestbook"))

# =========================================================
# EDIT BUKU TAMU
# =========================================================
# =========================================================
# EDIT BUKU TAMU
# =========================================================

@app.route("/buku-tamu/edit/<int:id>", methods=["POST"])
@login_required
def edit_guest(id):

    data = request.form

    db = get_db()
    cur = db.cursor()

    cur.execute("""
        UPDATE guestbook
        SET
            name=%s,
            origin=%s,
            purpose=%s,
            visit_date=%s,
            entry_time=%s,
            exit_time=%s,
            note=%s,
            paraf=%s
        WHERE id=%s
    """, (
        data["name"],
        data["origin"],
        data["purpose"],
        data["visit_date"],
        data["entry_time"],
        data.get("exit_time") or None,
        data.get("note", ""),
        data.get("paraf", ""),
        id
    ))

    db.commit()

    cur.close()
    db.close()

    flash(
        "Data tamu berhasil diperbarui.",
        "success"
    )

    return redirect(url_for("guestbook"))

# =========================================================
# HAPUS BUKU TAMU
# =========================================================

@app.route("/buku-tamu/hapus/<int:id>", methods=["POST"])
@login_required
def delete_guest(id):

    db = get_db()
    cur = db.cursor()


    cur.execute(
        "DELETE FROM guestbook WHERE id=%s",
        (id,)
    )


    db.commit()

    cur.close()
    db.close()


    flash(
        "Data tamu berhasil dihapus.",
        "success"
    )

    return redirect(url_for("guestbook"))


# =========================================================
# RUN
# =========================================================

if __name__ == "__main__":

    ensure_admin_account()

    app.run(
        debug=True
    )